"""MOCRE-1 F2/Oklahoma: ingest real OCC UIC injection volumes (the fluid-pressure variable, the
single highest-confidence finding of the 2026-07-03 catalog audit -- mu_ev=0.90, previously
marked "almost never public", which was false for this exact well-documented case).

Source: annual OCC spreadsheets, monthly Vol (bbls) + PSI per well, real lat/lon.
NOT a live API -- one xlsx per year, downloaded and cached locally, re-run yearly for new data.
"""
import os
import sys

import openpyxl
import pandas as pd

from common import ROOT, dist_to_trace_km, load_config, segment_paths

URL_TMPL = "https://oklahoma.gov/content/dam/ok/en/occ/documents/og/ogdatafiles/{year}-uic-injection-volumes.xlsx"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
PUBLICATION_MONTH = int(os.getenv("MOCRE_OCC_PUBLICATION_MONTH", "2"))
PUBLICATION_DAY = int(os.getenv("MOCRE_OCC_PUBLICATION_DAY", "15"))


def download_year(year, cache_dir):
    path = os.path.join(cache_dir, f"{year}.xlsx")
    if os.path.exists(path):
        return path
    import urllib.request
    url = URL_TMPL.format(year=year)
    print(f"  downloading {year}...")
    urllib.request.urlretrieve(url, path)
    return path


def parse_year(path, cfg):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {name: i for i, name in enumerate(header)}
    lat_i, lon_i = idx["LAT"], idx["LON"]

    lats, lons = [], []
    valid_rows = []
    for r in rows[1:]:
        lat, lon = r[lat_i], r[lon_i]
        if lat is None or lon is None:
            continue
        try:
            lat, lon = float(lat), float(lon)
        except (TypeError, ValueError):
            continue  # e.g. literal "NULL" string seen in some year files
        lats.append(lat)
        lons.append(lon)
        valid_rows.append(r)
    if not valid_rows:
        return []

    dist, _ = dist_to_trace_km(lats, lons, cfg)
    max_d = cfg["corridor_half_width_km"]
    out = []
    for r, d in zip(valid_rows, dist):
        if d > max_d:
            continue
        year = r[idx["ReportYear"]].year if r[idx["ReportYear"]] else None
        for m in MONTHS:
            vol = r[idx[f"{m} Vol"]]
            psi = r[idx[f"{m} PSI"]]
            if vol is None and psi is None:
                continue
            out.append({"year": year, "month": m, "api": r[idx["API"]],
                       "lat": r[lat_i], "lon": r[lon_i], "dist_km": round(float(d), 2),
                       "vol_bbls": vol, "psi": psi})
    return out


def main():
    cfg = load_config()
    paths = segment_paths(cfg)
    print(f"segment: {cfg['segment_id']} ({paths['slug']})")
    if cfg["regime"] != "induced":
        print("WARNING: this segment isn't marked regime=induced -- wells data may not apply.")

    wells_dir = os.path.join(ROOT, "data", "wells", paths["slug"])
    os.makedirs(wells_dir, exist_ok=True)
    cache_dir = os.path.join(ROOT, "data", "wells", "_xlsx_cache")
    os.makedirs(cache_dir, exist_ok=True)

    ya, yb = cfg["wells"]["years_available"].split("-")
    all_rows = []
    for year in range(int(ya), int(yb) + 1):
        try:
            path = download_year(year, cache_dir)
        except Exception as e:  # noqa: BLE001
            print(f"  {year}: download failed ({e}), skipping")
            continue
        rows = parse_year(path, cfg)
        print(f"  {year}: {len(rows)} well-month records in corridor")
        all_rows.extend(rows)

    df = pd.DataFrame(all_rows)
    out_path = os.path.join(wells_dir, "well_months_raw.csv")
    df.to_csv(out_path, index=False)
    print(f"wrote {len(df)} well-month records -> {out_path}")

    # Monthly aggregate across all wells in corridor: sum(vol), vol-weighted mean(psi).
    month_num = {m: i + 1 for i, m in enumerate(MONTHS)}
    df["month_num"] = df["month"].map(month_num)
    df["date"] = pd.to_datetime(dict(year=df["year"], month=df["month_num"], day=1))
    # No existe un archivo versionado que demuestre cuándo estuvo disponible
    # cada fila mensual. Para backtest se aplica el supuesto conservador de
    # publicación del libro anual el 15-feb del año siguiente.
    df["available_at"] = pd.to_datetime(dict(year=df["year"] + 1,
                                              month=PUBLICATION_MONTH,
                                              day=PUBLICATION_DAY))
    df["vol_bbls"] = pd.to_numeric(df["vol_bbls"], errors="coerce").fillna(0)
    df["psi"] = pd.to_numeric(df["psi"], errors="coerce")

    def agg(g):
        total_vol = g["vol_bbls"].sum()
        w = g["vol_bbls"].to_numpy()
        p = g["psi"].to_numpy()
        mask = ~pd.isna(p) & (w > 0)
        psi_w = (p[mask] * w[mask]).sum() / w[mask].sum() if mask.sum() > 0 else float("nan")
        return pd.Series({"total_vol_bbls": total_vol, "vol_weighted_psi": psi_w,
                          "n_wells": g["api"].nunique()})

    monthly = df.groupby("date").apply(agg, include_groups=False).sort_index()
    monthly["available_at"] = [pd.Timestamp(year=d.year + 1, month=PUBLICATION_MONTH,
                                             day=PUBLICATION_DAY) for d in monthly.index]
    monthly_path = os.path.join(wells_dir, "monthly_aggregate.csv")
    monthly.to_csv(monthly_path)
    print(f"wrote {len(monthly)} monthly rows -> {monthly_path}")
    print(monthly.tail(10))


if __name__ == "__main__":
    main()
